from pathlib import Path
import json, math, hashlib, zipfile, csv, collections
from openpyxl import load_workbook
root=Path.cwd(); out=root/'output/木质素碳点训练数据库'
data=json.loads((root/'tmp/database.json').read_text(encoding='utf8'))
book=out/'PET搭配材料筛选_木质素碳点文献数据库.xlsx'
wb=load_workbook(book,data_only=False)
assert list(wb.sheetnames)==list(data['tables'])
cells=0
for name,records in data['tables'].items():
    sh=wb[name]; headers=data['headers'][name]
    assert [sh.cell(1,j+1).value for j in range(len(headers))]==headers
    assert not sh.merged_cells.ranges
    for i,rec in enumerate(records,2):
        for j,h in enumerate(headers,1):
            expected=rec.get(h); cell=sh.cell(i,j); actual=cell.value; cells+=1
            assert cell.data_type not in ['f','e'],(name,cell.coordinate,cell.data_type,actual)
            if isinstance(expected,(float,int)):
                assert isinstance(actual,(float,int)) and math.isclose(expected,actual,rel_tol=1e-12,abs_tol=1e-12),(name,cell.coordinate,expected,actual)
            else: assert expected==actual or expected is None and actual is None,(name,cell.coordinate,expected,actual)
    if records: assert len(sh.tables)==1
    if name=='PET实验录入': assert all(c.value is None for row in sh.iter_rows(min_row=2) for c in row)
    print(name,'rows',len(records),'freeze',sh.freeze_panes)
src=root/'文献-木质素基抗菌碳点'/'文献-木质素基抗菌碳点'
for rec in json.loads((root/'tmp/extracted/manifest.json').read_text(encoding='utf8')):
    assert hashlib.sha256((src/rec['file'].replace('\\','/')).read_bytes()).hexdigest()==rec['sha256']
train=data['tables']['抗菌训练长表']; free=[r for r in train if r['material_context']=='free_CD' and r['regression_eligible']==1 and r['data_status']!='derived']
print('free_CD exact numeric records',len(free),'unique CD samples',len({r['sample_id'] for r in free}))
assert len(free)==6
print('training use',dict(collections.Counter(r['default_training_use'] for r in train)))
with zipfile.ZipFile(out/'机器学习CSV数据包.zip') as z:
    assert len(z.namelist())==len(data['tables'])+1
    assert z.testzip() is None
print('verified',cells,'data cells; originals unchanged; no Excel formulas/errors; CSV ZIP valid')
(root/'tmp/verification_result.json').write_text(json.dumps({'checked_cells':cells,'sheet_count':len(wb.sheetnames),'free_CD_exact_numeric_records':len(free),'free_CD_unique_samples':len({r['sample_id'] for r in free}),'source_hashes_unchanged':True},indent=2),encoding='utf8')
